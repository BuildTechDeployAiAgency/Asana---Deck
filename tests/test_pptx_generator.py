"""Tests for PowerPoint generator."""

import tempfile
from pathlib import Path

from asana_deck.config import Config
from asana_deck.models import DeckContent, SlideContent
from asana_deck.generators.pptx_generator import PptxGenerator


def _sample_deck() -> DeckContent:
    return DeckContent(
        title="Test Deck",
        subtitle="Unit tests",
        date="2026-02-18",
        slides=[
            SlideContent(
                title="Section One",
                subtitle="3 items",
                slide_type="section",
                section_name="Section One",
            ),
            SlideContent(
                title="Feature Alpha",
                subtitle="Core feature",
                bullets=["Fast", "Reliable", "Secure"],
                executive_summary="Alpha is the cornerstone feature.",
                slide_type="content",
                speaker_notes="Mention the benchmark results.",
                section_name="Section One",
            ),
            SlideContent(
                title="Feature Beta",
                bullets=["Easy to use", "Well documented"],
                slide_type="content",
                section_name="Section One",
            ),
        ],
    )


def test_generate_creates_file():
    config = Config(
        asana_access_token="x",
        anthropic_api_key="x",
    )
    gen = PptxGenerator(config)
    deck = _sample_deck()

    with tempfile.TemporaryDirectory() as tmpdir:
        out = Path(tmpdir) / "test_output.pptx"
        result = gen.generate(deck, out)
        assert result == out
        assert out.exists()
        assert out.stat().st_size > 0


def test_generate_creates_parent_dirs():
    config = Config(
        asana_access_token="x",
        anthropic_api_key="x",
    )
    gen = PptxGenerator(config)
    deck = _sample_deck()

    with tempfile.TemporaryDirectory() as tmpdir:
        out = Path(tmpdir) / "nested" / "dir" / "deck.pptx"
        result = gen.generate(deck, out)
        assert result.exists()


def test_slide_count():
    """The generated deck should have 1 title + N content slides."""
    config = Config(
        asana_access_token="x",
        anthropic_api_key="x",
    )
    gen = PptxGenerator(config)
    deck = _sample_deck()

    with tempfile.TemporaryDirectory() as tmpdir:
        out = Path(tmpdir) / "count_test.pptx"
        gen.generate(deck, out)

        from pptx import Presentation

        prs = Presentation(str(out))
        # 1 title slide + 1 section slide + 2 content slides = 4
        assert len(prs.slides) == 4


def test_speaker_notes():
    config = Config(asana_access_token="x", anthropic_api_key="x")
    gen = PptxGenerator(config)
    deck = _sample_deck()

    with tempfile.TemporaryDirectory() as tmpdir:
        out = Path(tmpdir) / "notes_test.pptx"
        gen.generate(deck, out)

        from pptx import Presentation

        prs = Presentation(str(out))
        # Slide index 2 is the first content slide with speaker notes
        notes = prs.slides[2].notes_slide.notes_text_frame.text
        assert "benchmark" in notes.lower()


def test_widescreen_dimensions():
    config = Config(asana_access_token="x", anthropic_api_key="x")
    gen = PptxGenerator(config)
    deck = DeckContent(title="Minimal", slides=[])

    with tempfile.TemporaryDirectory() as tmpdir:
        out = Path(tmpdir) / "dim_test.pptx"
        gen.generate(deck, out)

        from pptx import Presentation

        prs = Presentation(str(out))
        width_inches = prs.slide_width / 914400
        assert abs(width_inches - 13.333) < 0.01
