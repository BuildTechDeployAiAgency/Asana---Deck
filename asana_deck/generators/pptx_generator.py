"""PowerPoint generator — builds .pptx files from DeckContent using python-pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from asana_deck.config import Config
from asana_deck.models import DeckContent, SlideContent


def _hex_to_rgb(hex_color: str) -> RGBColor:
    """Convert '#RRGGBB' to an RGBColor."""
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class PptxGenerator:
    """Generates a PowerPoint deck from structured DeckContent."""

    def __init__(self, config: Config) -> None:
        self._config = config
        self._prs = Presentation()
        # Widescreen 16:9
        self._prs.slide_width = Inches(13.333)
        self._prs.slide_height = Inches(7.5)

    # ── Colour / font helpers ─────────────────────────────────────────

    @property
    def _primary(self) -> RGBColor:
        return _hex_to_rgb(self._config.brand_color_primary)

    @property
    def _accent(self) -> RGBColor:
        return _hex_to_rgb(self._config.brand_color_accent)

    @property
    def _text_color(self) -> RGBColor:
        return _hex_to_rgb(self._config.brand_color_text)

    def _set_font(
        self,
        run,
        *,
        size: int = 18,
        bold: bool = False,
        color: RGBColor | None = None,
        font_name: str | None = None,
    ) -> None:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color or self._text_color
        run.font.name = font_name or self._config.brand_font_body

    # ── Background helper ─────────────────────────────────────────────

    def _set_slide_bg(self, slide, color: RGBColor) -> None:
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # ── Slide builders ────────────────────────────────────────────────

    def _add_title_slide(self, deck: DeckContent) -> None:
        slide = self._prs.slides.add_slide(self._prs.slide_layouts[6])  # blank
        self._set_slide_bg(slide, self._primary)

        # Title
        left, top = Inches(1), Inches(2.2)
        txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = deck.title
        self._set_font(
            run,
            size=44,
            bold=True,
            font_name=self._config.brand_font_title,
        )

        # Subtitle
        if deck.subtitle:
            txBox2 = slide.shapes.add_textbox(
                Inches(1), Inches(4), Inches(11), Inches(1)
            )
            tf2 = txBox2.text_frame
            tf2.word_wrap = True
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.CENTER
            run2 = p2.add_run()
            run2.text = deck.subtitle
            self._set_font(run2, size=24, color=self._accent)

        # Date
        if deck.date:
            txBox3 = slide.shapes.add_textbox(
                Inches(1), Inches(5.2), Inches(11), Inches(0.6)
            )
            tf3 = txBox3.text_frame
            p3 = tf3.paragraphs[0]
            p3.alignment = PP_ALIGN.CENTER
            run3 = p3.add_run()
            run3.text = deck.date
            self._set_font(run3, size=16)

        # Accent bar
        slide.shapes.add_shape(
            1,  # rectangle
            Inches(4.5),
            Inches(3.8),
            Inches(4),
            Inches(0.06),
        ).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = self._accent

    def _add_section_slide(self, content: SlideContent) -> None:
        slide = self._prs.slides.add_slide(self._prs.slide_layouts[6])
        self._set_slide_bg(slide, self._accent)

        txBox = slide.shapes.add_textbox(
            Inches(1), Inches(2.5), Inches(11), Inches(1.5)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = content.title
        self._set_font(
            run,
            size=40,
            bold=True,
            font_name=self._config.brand_font_title,
        )

        if content.subtitle:
            txBox2 = slide.shapes.add_textbox(
                Inches(1), Inches(4.2), Inches(11), Inches(0.8)
            )
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.CENTER
            run2 = p2.add_run()
            run2.text = content.subtitle
            self._set_font(run2, size=20)

    def _add_content_slide(self, content: SlideContent) -> None:
        slide = self._prs.slides.add_slide(self._prs.slide_layouts[6])
        self._set_slide_bg(slide, _hex_to_rgb("#FFFFFF"))

        # Left accent bar
        bar = slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(0.15), Inches(7.5)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = self._accent
        bar.line.fill.background()

        # Title
        txBox = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.5), Inches(11.5), Inches(1)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = content.title
        self._set_font(
            run,
            size=32,
            bold=True,
            color=self._primary,
            font_name=self._config.brand_font_title,
        )

        # Subtitle
        if content.subtitle:
            txBox2 = slide.shapes.add_textbox(
                Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.6)
            )
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            run2 = p2.add_run()
            run2.text = content.subtitle
            self._set_font(run2, size=18, color=self._accent)

        # Bullets
        if content.bullets:
            bullet_top = Inches(2.3) if content.subtitle else Inches(1.8)
            txBox3 = slide.shapes.add_textbox(
                Inches(1.2), bullet_top, Inches(10.5), Inches(4)
            )
            tf3 = txBox3.text_frame
            tf3.word_wrap = True
            for i, bullet in enumerate(content.bullets):
                p3 = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
                p3.space_after = Pt(10)
                run3 = p3.add_run()
                run3.text = f"\u2022  {bullet}"
                self._set_font(run3, size=18, color=_hex_to_rgb("#333333"))

        # Executive summary at bottom
        if content.executive_summary:
            txBox4 = slide.shapes.add_textbox(
                Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.8)
            )
            tf4 = txBox4.text_frame
            tf4.word_wrap = True
            p4 = tf4.paragraphs[0]
            run4 = p4.add_run()
            run4.text = content.executive_summary
            self._set_font(
                run4, size=14, bold=True, color=self._accent
            )

        # Speaker notes
        if content.speaker_notes:
            slide.notes_slide.notes_text_frame.text = content.speaker_notes

    # ── Public API ────────────────────────────────────────────────────

    def generate(self, deck: DeckContent, output_path: Path | str) -> Path:
        """Render a full deck and save to *output_path*. Returns the path."""
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Title slide
        self._add_title_slide(deck)

        # Content slides
        for slide_content in deck.slides:
            if slide_content.slide_type == "section":
                self._add_section_slide(slide_content)
            else:
                self._add_content_slide(slide_content)

        self._prs.save(str(output_path))
        return output_path
